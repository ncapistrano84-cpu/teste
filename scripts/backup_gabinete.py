"""Gera o backup mensal (Excel) do Painel de Gestão do Gabinete.

Uso:
    python3 scripts/backup_gabinete.py <pasta_com_dados> <arquivo_saida.xlsx>

<pasta_com_dados> é a pasta onde a ferramenta ArtifactData salvou os documentos
(out_dir), com uma subpasta por coleção: gab_lotes/, gab_notas/, gab_sentencas/,
gab_mov/ e gab_config/equipe.json. Cada arquivo .json contém os campos do
documento; o id é o nome do arquivo.
"""
import datetime
import glob
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

STN = {"p": "Pendente", "c": "Concluído", "d": "Dúvida", "e": "Erro"}
MC = {"remessas": "Remessas e peritos", "extincao": "Extinção – arquivamento",
      "abandono": "Intimar sob pena de abandono", "arqsembaixa": "Arquivar sem baixa"}
SS = {"pendente": "Aguardando", "analise": "Em análise", "assinada": "Assinada"}
ST2 = {"merito": "Mérito", "extincao": "Extinção"}


def carregar(pasta, colecao):
    docs = []
    for f in sorted(glob.glob(os.path.join(pasta, colecao, "*.json"))):
        with open(f, encoding="utf-8") as fh:
            d = json.load(fh)
        d = d.get("data", d) if isinstance(d, dict) and "data" in d and isinstance(d["data"], dict) else d
        d.setdefault("id", os.path.splitext(os.path.basename(f))[0])
        docs.append(d)
    return docs


def data_br(v):
    if not v:
        return ""
    try:
        if len(v) == 10:
            return datetime.date.fromisoformat(v).strftime("%d/%m/%Y")
        dt = datetime.datetime.fromisoformat(v.replace("Z", "+00:00")) - datetime.timedelta(hours=3)
        return dt.strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return str(v)


def metas(l):
    if isinstance(l.get("metas"), list):
        return ", ".join(m for m in l["metas"] if m)
    return l.get("meta") or ""


def main(pasta, saida):
    lotes = sorted(carregar(pasta, "gab_lotes"), key=lambda l: l.get("criadoEm") or "")
    notas = sorted(carregar(pasta, "gab_notas"), key=lambda n: n.get("criadoEm") or "")
    sents = sorted(carregar(pasta, "gab_sentencas"), key=lambda s: s.get("conclusao") or "")
    movs = carregar(pasta, "gab_mov")
    eq_path = os.path.join(pasta, "gab_config", "equipe.json")
    equipe = []
    if os.path.exists(eq_path):
        with open(eq_path, encoding="utf-8") as fh:
            e = json.load(fh)
        equipe = (e.get("data", e)).get("membros", [])
    grupo = {m.get("nome"): m.get("grupo", "") for m in equipe}

    proc = [["Serviço (id)", "Tipo", "Subtipo", "Metas", "Executor", "Distribuído por", "Distribuído em", "Prazo", "Arquivado",
             "Nº do processo", "Descrição", "Situação", "Marcado por", "Marcado em"]]
    serv = [["Serviço (id)", "Tipo", "Subtipo", "Metas", "Executor", "Distribuído por", "Distribuído em", "Prazo", "Instruções",
             "Origem", "Arquivado em", "Total", "Concluídos", "Pendentes", "Dúvidas", "Erros", "% concluído"]]
    tot = dict(t=0, c=0, p=0, d=0, e=0)
    por_ex = {}
    for l in lotes:
        c = dict(t=0, c=0, p=0, d=0, e=0)
        st = l.get("st") or {}
        for i, p in enumerate(l.get("procs") or []):
            x = st.get(f"k{i}") or {}
            s = x.get("s") or "p"
            c["t"] += 1
            c[s] += 1
            proc.append([l["id"], l.get("tipo", ""), l.get("subtipo", ""), metas(l), l.get("executor", ""), l.get("criadoPor", ""),
                         data_br(l.get("criadoEm")), data_br(l.get("prazo")), "Sim" if l.get("arquivado") else "Não",
                         p.get("n", ""), p.get("i", ""), STN.get(s, s), x.get("by", ""), data_br(x.get("at"))])
        arq = l.get("arquivado") or {}
        serv.append([l["id"], l.get("tipo", ""), l.get("subtipo", ""), metas(l), l.get("executor", ""), l.get("criadoPor", ""),
                     data_br(l.get("criadoEm")), data_br(l.get("prazo")), l.get("obs", ""), l.get("origem", ""),
                     data_br(arq.get("em")) if arq else "", c["t"], c["c"], c["p"], c["d"], c["e"],
                     round(c["c"] / c["t"] * 100, 1) if c["t"] else 0])
        if not arq:
            for k in tot:
                tot[k] += c[k]
            r = por_ex.setdefault(l.get("executor", ""), dict(t=0, c=0, p=0, d=0, e=0))
            for k in r:
                r[k] += c[k]
    exe = [["Executor", "Grupo", "Total", "Concluídos", "Pendentes", "Dúvidas", "Erros", "% concluído"]] + [
        [n, grupo.get(n, ""), r["t"], r["c"], r["p"], r["d"], r["e"], round(r["c"] / r["t"] * 100, 1) if r["t"] else 0]
        for n, r in por_ex.items()]

    sen = [["Nº do processo", "Sistema", "Tipo", "Data da conclusão", "Situação", "Assinada/analisada em", "Por", "Cadastrado por",
            "Cadastrado em", "Observação"]] + [
        [x.get("n", ""), x.get("sistema", ""), ST2.get(x.get("tipo"), ""), data_br(x.get("conclusao")), SS.get(x.get("status"), x.get("status", "")),
         data_br(x.get("statusEm")), x.get("statusPor", ""), x.get("cadastradoPor", ""), data_br(x.get("cadastradoEm")), x.get("obs", "")]
        for x in sents]
    pm = {}
    for x in sents:
        if x.get("status") == "assinada" and x.get("statusEm"):
            r = pm.setdefault(x["statusEm"][:7], [0, 0])
            r[1 if x.get("tipo") == "extincao" else 0] += 1
    sen_m = [["Mês", "Mérito", "Extinção", "Total"]] + [[k[5:] + "/" + k[:4], v[0], v[1], v[0] + v[1]] for k, v in sorted(pm.items(), reverse=True)]

    mov = [["Nº do processo", "Descrição", "Sistema", "Finalidade", "Motivo", "Último andamento", "Prazo (dias)", "Data de averiguação",
            "Situação", "Observação", "Responsável", "Status", "Concluído em", "Concluído por", "Cadastrado por", "Cadastrado em"]] + [
        [m.get("n", ""), m.get("info", ""), m.get("sistema", ""), MC.get(m.get("categoria"), "Não informada"), m.get("motivo", ""),
         data_br(m.get("ultimo")), "" if m.get("prazo") is None else m.get("prazo"), data_br(m.get("averiguacao")), m.get("situacao", ""),
         m.get("obs", ""), m.get("responsavel", ""), "Concluído/Arquivado" if m.get("status") == "concluido" else "Em movimentação",
         data_br(m.get("concluidoEm")), m.get("concluidoPor", ""), m.get("cadastradoPor", ""), data_br(m.get("cadastradoEm"))]
        for m in movs]
    con = [["Data", "Autor", "Mencionados", "Mensagem", "Resposta à mensagem (id)", "Serviço", "Processo", "Lida por", "Id"]] + [
        [data_br(n.get("criadoEm")), n.get("autor", ""), ", ".join(n.get("mencoes") or []), n.get("texto", ""), n.get("thread") or "",
         n.get("loteDesc") or "", n.get("processo") or "", ", ".join((n.get("lidos") or {}).keys()), n["id"]] for n in notas]
    eq = [["Nome", "Grupo", "E-mail"]] + [[m.get("nome", ""), m.get("grupo", ""), m.get("email", "")] for m in equipe]
    agora = (datetime.datetime.now(datetime.timezone.utc) - datetime.timedelta(hours=3)).strftime("%d/%m/%Y %H:%M")
    res = [["Item", "Quantidade"], ["Backup gerado em", agora], ["Gerado por", "Rotina mensal automática"],
           ["Serviços distribuídos (todos)", len(lotes)], ["Serviços arquivados", sum(1 for l in lotes if l.get("arquivado"))],
           ["Processos distribuídos (ativos)", tot["t"]], ["Concluídos", tot["c"]], ["Pendentes", tot["p"]], ["Com dúvida", tot["d"]], ["Com erro", tot["e"]],
           ["Sentenças aguardando", sum(1 for x in sents if x.get("status") == "pendente")], ["Sentenças em análise", sum(1 for x in sents if x.get("status") == "analise")],
           ["Sentenças assinadas", sum(1 for x in sents if x.get("status") == "assinada")],
           ["Movimentação – em andamento", sum(1 for m in movs if m.get("status") != "concluido")], ["Movimentação – concluídos/arquivados", sum(1 for m in movs if m.get("status") == "concluido")],
           ["Mensagens (conversas)", len(notas)], ["Pessoas na equipe", len(equipe)]]

    wb = Workbook()
    wb.remove(wb.active)
    cab = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="4A6B5D")
    for nome, rows in [("Resumo", res), ("Processos distribuídos", proc), ("Serviços", serv), ("Por executor", exe), ("Sentenças", sen),
                       ("Sentenças por mês", sen_m), ("Movimentação", mov), ("Conversas", con), ("Equipe", eq)]:
        ws = wb.create_sheet(nome)
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = cab
            cell.fill = fundo
        ws.freeze_panes = "A2"
        if len(rows) > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"
        for i in range(len(rows[0])):
            larg = max(len(str(r[i])) if i < len(r) and r[i] is not None else 0 for r in rows[:300]) + 2
            ws.column_dimensions[get_column_letter(i + 1)].width = min(60, max(10, larg))
    wb.save(saida)
    print(f"Backup salvo em {saida}: {len(lotes)} serviços, {tot['t']} processos ativos, {len(sents)} sentenças, {len(movs)} movimentações, {len(notas)} mensagens.")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
